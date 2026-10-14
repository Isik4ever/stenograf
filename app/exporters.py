"""Exportformat: TXT, TXT med tider, Markdown, SRT, VTT, DOCX, JSON."""

from __future__ import annotations

import io
import json


def ts(sec: float, sep: str = ",", hours: bool = True) -> str:
    ms = int(round(max(0.0, sec) * 1000))
    h, ms = divmod(ms, 3_600_000)
    m, ms = divmod(ms, 60_000)
    s, ms = divmod(ms, 1000)
    if hours:
        return f"{h:02d}:{m:02d}:{s:02d}{sep}{ms:03d}"
    return f"{m:02d}:{s:02d}{sep}{ms:03d}"


def clock(sec: float) -> str:
    sec = int(max(0, sec))
    h, rem = divmod(sec, 3600)
    m, s = divmod(rem, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


def _paras(doc: dict) -> list[dict]:
    return doc["transcript"]["paragraphs"]


def _ptext(p: dict) -> str:
    return " ".join(s["text"] for s in p["sentences"])


def _names(doc: dict) -> dict[str, str]:
    from .textlayout import speaker_names

    return speaker_names(doc["transcript"], doc.get("speakers"))


def _who(doc: dict, p: dict, names: dict | None = None) -> str | None:
    names = _names(doc) if names is None else names
    spk = p.get("speaker")
    return names.get(str(spk)) if names and spk is not None else None


def _header(doc: dict) -> list[str]:
    meta = doc.get("meta", {})
    lines = [meta.get("title") or "Transkription"]
    bits = [b for b in (meta.get("uploader"), meta.get("url")) if b]
    if bits:
        lines.append(" · ".join(bits))
    return lines


def to_txt(doc: dict, timestamps: bool = False) -> str:
    names = _names(doc)
    out = []
    for p in _paras(doc):
        who = _who(doc, p, names)
        line = f"{who}: {_ptext(p)}" if who else _ptext(p)
        out.append(f"[{clock(p['start'])}] {line}" if timestamps else line)
    return "\n\n".join(out) + "\n"


def to_md(doc: dict) -> str:
    h = _header(doc)
    out = [f"# {h[0]}"]
    if len(h) > 1:
        out.append(f"_{h[1]}_")
    names = _names(doc)
    for p in _paras(doc):
        who = _who(doc, p, names)
        out.append(f"**{clock(p['start'])}** " + (f"**{who}:** " if who else "— ") + _ptext(p))
    return "\n\n".join(out) + "\n"


def to_srt(doc: dict) -> str:
    out = []
    for i, c in enumerate(doc["transcript"]["cues"], 1):
        out.append(f"{i}\n{ts(c['start'])} --> {ts(c['end'])}\n{c['text']}\n")
    return "\n".join(out)


def to_vtt(doc: dict) -> str:
    out = ["WEBVTT", ""]
    for c in doc["transcript"]["cues"]:
        out.append(f"{ts(c['start'], '.')} --> {ts(c['end'], '.')}\n{c['text']}\n")
    return "\n".join(out)


def to_json(doc: dict) -> str:
    return json.dumps(doc, ensure_ascii=False, indent=2)


def to_docx(doc: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Pt, RGBColor

    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    h = _header(doc)
    d.add_heading(h[0], level=1)
    if len(h) > 1:
        sub = d.add_paragraph(h[1])
        sub.runs[0].font.color.rgb = RGBColor(0x6B, 0x6B, 0x6B)
        sub.runs[0].font.size = Pt(9)
    names = _names(doc)
    for p in _paras(doc):
        para = d.add_paragraph()
        t = para.add_run(clock(p["start"]) + "  ")
        t.font.size = Pt(8.5)
        t.font.color.rgb = RGBColor(0x8A, 0x8A, 0x8A)
        t.font.highlight_color = WD_COLOR_INDEX.GRAY_25
        who = _who(doc, p, names)
        if who:
            para.add_run(who + ": ").bold = True
        para.add_run(_ptext(p))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


FORMATS = {
    "txt": ("text/plain; charset=utf-8", lambda d: to_txt(d)),
    "timed.txt": ("text/plain; charset=utf-8", lambda d: to_txt(d, timestamps=True)),
    "md": ("text/markdown; charset=utf-8", to_md),
    "srt": ("application/x-subrip; charset=utf-8", to_srt),
    "vtt": ("text/vtt; charset=utf-8", to_vtt),
    "json": ("application/json; charset=utf-8", to_json),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", to_docx),
}
